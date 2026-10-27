import base64
import json
import os
import shutil
import subprocess
import time
import urllib.request
import websocket

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCREENSHOT_DIR = os.path.join(BASE_DIR, "docs", "screenshots")
STATIC_DIR = os.path.join(BASE_DIR, "static")
DOWNLOADS_DIR = r"C:\Users\2427727\Downloads"
ARTIFACT_DIR = r"C:\Users\2427727\.gemini\antigravity\brain\28d738c2-5a3d-47f9-907e-bddbfc7bf8c1"
os.makedirs(SCREENSHOT_DIR, exist_ok=True)
os.makedirs(ARTIFACT_DIR, exist_ok=True)

CHROME_PATHS = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
]


def get_browser_path():
    for p in CHROME_PATHS:
        if os.path.exists(p):
            return p
    raise RuntimeError("Neither Chrome nor Edge found.")


class CDPBrowser:
    def __init__(self, port=9234):
        self.port = port
        self.proc = None
        self.ws = None
        self.msg_id = 0

    def start(self):
        browser_exe = get_browser_path()
        user_data_dir = os.path.join(BASE_DIR, "data", "chrome_ppt_profile")
        os.makedirs(user_data_dir, exist_ok=True)
        cmd = [
            browser_exe,
            "--headless=new",
            f"--remote-debugging-port={self.port}",
            "--remote-allow-origins=*",
            "--window-size=1600,960",
            "--hide-scrollbars",
            "--disable-gpu",
            "--no-first-run",
            "--no-default-browser-check",
            f"--user-data-dir={user_data_dir}",
            "http://127.0.0.1:8050/",
        ]
        self.proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(30):
            try:
                resp = urllib.request.urlopen(f"http://127.0.0.1:{self.port}/json").read().decode()
                targets = json.loads(resp)
                page_target = next((t for t in targets if t.get("type") == "page"), None)
                if page_target and "webSocketDebuggerUrl" in page_target:
                    self.ws = websocket.create_connection(page_target["webSocketDebuggerUrl"])
                    break
            except Exception:
                time.sleep(0.3)
        if not self.ws:
            raise RuntimeError("Failed to connect to headless browser via CDP.")

        self.send("Page.enable")
        self.send("Runtime.enable")
        self.send("Emulation.setDeviceMetricsOverride", {
            "width": 1600,
            "height": 960,
            "deviceScaleFactor": 2,
            "mobile": False,
        })

    def send(self, method, params=None):
        self.msg_id += 1
        payload = {"id": self.msg_id, "method": method, "params": params or {}}
        self.ws.send(json.dumps(payload))
        while True:
            raw = self.ws.recv()
            msg = json.loads(raw)
            if msg.get("id") == self.msg_id:
                return msg.get("result", {})

    def navigate(self, url, wait_sec=1.5):
        self.send("Page.navigate", {"url": url})
        time.sleep(wait_sec)

    def evaluate(self, expression, wait_sec=0.6):
        res = self.send("Runtime.evaluate", {
            "expression": expression,
            "returnByValue": True,
            "awaitPromise": True,
        })
        if wait_sec > 0:
            time.sleep(wait_sec)
        return res

    def screenshot(self, filename, clip=None):
        params = {"format": "png", "fromSurface": True}
        if clip:
            params["clip"] = clip
        res = self.send("Page.captureScreenshot", params)
        data = base64.b64decode(res["data"])
        out_path = os.path.join(SCREENSHOT_DIR, filename)
        with open(out_path, "wb") as f:
            f.write(data)
        art_path = os.path.join(ARTIFACT_DIR, filename)
        with open(art_path, "wb") as f:
            f.write(data)
        print(f"Saved screenshot: {out_path}")
        return out_path

    def close(self):
        try:
            if self.ws:
                self.ws.close()
        except Exception:
            pass
        try:
            if self.proc:
                self.proc.terminate()
        except Exception:
            pass


def reset_app_state():
    req = urllib.request.Request("http://127.0.0.1:8050/api/reset", method="POST")
    urllib.request.urlopen(req).read()


def capture_all_product_screenshots():
    print("Resetting databases to clean synthetic baseline...")
    reset_app_state()

    browser = CDPBrowser()
    try:
        browser.start()
        # 1. Main Control Tower Overview (Top KPIs + 5-Agent Strip + 24-Hr Gantt Chart - Full Mix)
        browser.navigate("http://127.0.0.1:8050/", wait_sec=2.0)
        browser.evaluate("scrollGanttToHour(7.8, false); window.scrollTo(0, 0);", wait_sec=0.8)
        browser.screenshot("slide_01_control_tower_overview.png")

        # 2. Focused 24-Hour Scrollable Gantt Chart (Frozen Left Labels + 30/45/60m Bars + Red Overlap Indicator)
        browser.evaluate("""
            const panel = document.getElementById('tabpanel-gantt');
            if (panel) panel.scrollIntoView({ behavior: 'instant', block: 'start' });
            window.scrollBy(0, -10);
            scrollGanttToHour(7.85, false);
        """, wait_sec=0.8)
        browser.screenshot("slide_02_gantt_24h_frozen_and_overlaps.png")

        # 3. SCENARIO 1: PAIRWISE SLOT SWAP (SCENARIO_SLOT_SWAP) — Gantt View
        browser.evaluate("""
            switchViewTab('gantt');
            loadScenario('SCENARIO_SLOT_SWAP', document.getElementById('scen-slot'));
        """, wait_sec=1.2)
        browser.evaluate("""
            const banner = document.getElementById('toastBanner');
            if (banner) banner.hidden = true;
            window.scrollTo(0, 0);
            scrollGanttToHour(7.75, false);
        """, wait_sec=0.6)
        browser.screenshot("slide_03_slot_swap_gantt.png")

        # 4. SCENARIO 1: PAIRWISE SLOT SWAP — Fleet Table & Agent Proposal Card View
        browser.evaluate("""
            switchViewTab('ops');
            window.scrollTo(0, 150);
        """, wait_sec=0.8)
        browser.screenshot("slide_04_slot_swap_proposals.png")

        # 5. SCENARIO 2: CROSS-DOCK SWAP (SCENARIO_DOCK_SWAP: DOCK-A <-> DOCK-B) — Gantt View
        browser.evaluate("""
            switchViewTab('gantt');
            loadScenario('SCENARIO_DOCK_SWAP', document.getElementById('scen-dock'));
        """, wait_sec=1.2)
        browser.evaluate("""
            const banner = document.getElementById('toastBanner');
            if (banner) banner.hidden = true;
            window.scrollTo(0, 0);
            scrollGanttToHour(7.85, false);
        """, wait_sec=0.6)
        browser.screenshot("slide_05_dock_swap_gantt.png")

        # 6. SCENARIO 2: CROSS-DOCK SWAP — Fleet Table & Agent Proposal Card View
        browser.evaluate("""
            switchViewTab('ops');
            window.scrollTo(0, 150);
        """, wait_sec=0.8)
        browser.screenshot("slide_06_dock_swap_proposals.png")

        # Restore Full Multi-Agent Scenario for remaining screenshots
        browser.evaluate("""
            switchViewTab('gantt');
            loadScenario('SCENARIO_FULL_MULTI_AGENT', document.getElementById('scen-full'));
        """, wait_sec=1.2)
        browser.evaluate("""
            const banner = document.getElementById('toastBanner');
            if (banner) banner.hidden = true;
            const panel = document.getElementById('tabpanel-gantt');
            if (panel) panel.scrollIntoView({ behavior: 'instant', block: 'start' });
            window.scrollBy(0, -10);
            scrollGanttToHour(7.85, false);
        """, wait_sec=0.6)

        # 7. Custom ETA Override Modal Dialog on Gantt Chart
        browser.evaluate("openETAModal('TRK-101');", wait_sec=0.7)
        browser.screenshot("slide_07_eta_modal_override.png")
        browser.evaluate("document.getElementById('etaModal').close();", wait_sec=0.4)

        # 8. Tab 2: Truck Fleet, ETA Controls & Logs (Full Mix)
        browser.evaluate("switchViewTab('ops'); window.scrollTo(0, 175);", wait_sec=0.8)
        browser.screenshot("slide_08_fleet_table_and_logs.png")

        # 9. Tab 3: DB & External System Connections Hub (Truck DB, Schedule DB, Live ETA Telematics)
        browser.evaluate("switchViewTab('connections'); window.scrollTo(0, 175);", wait_sec=0.8)
        browser.screenshot("slide_09_db_and_external_connections.png")

        # 10. Standalone Transport Portal (/transport) — 1. Manual ETA Input & 2. Book New Truck Schedule (30/45/60m)
        browser.navigate("http://127.0.0.1:8050/transport", wait_sec=1.8)
        browser.evaluate("window.scrollTo(0, 0);", wait_sec=0.6)
        browser.screenshot("slide_10_transport_portal.png")

        # 11. Approved & Committed Schedule State on 24-Hour Gantt Chart
        browser.navigate("http://127.0.0.1:8050/", wait_sec=1.8)
        browser.evaluate("executeMultiAgentPipeline();", wait_sec=1.5)
        browser.evaluate("""
            const banner = document.getElementById('toastBanner');
            if (banner) banner.hidden = true;
            const panel = document.getElementById('tabpanel-gantt');
            if (panel) panel.scrollIntoView({ behavior: 'instant', block: 'start' });
            window.scrollBy(0, -10);
            scrollGanttToHour(7.85, false);
        """, wait_sec=0.8)
        browser.screenshot("slide_11_approved_committed_schedule.png")

    finally:
        browser.close()
        # Restore baseline state so the user's live app is ready with proposals
        reset_app_state()


def add_header_banner(slide, prs, step_badge, title_text, subtitle_text):
    header_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.12)
    )
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = RGBColor(15, 23, 42)  # #0f172a
    header_box.line.fill.background()

    accent_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.08), Inches(13.333), Inches(0.05)
    )
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = RGBColor(37, 99, 235)  # #2563eb
    accent_line.line.fill.background()

    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.22), Inches(1.38), Inches(0.32)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(37, 99, 235)
    pill.line.fill.background()
    tf_p = pill.text_frame
    tf_p.word_wrap = False
    tf_p.margin_left = tf_p.margin_right = tf_p.margin_top = tf_p.margin_bottom = 0
    p_para = tf_p.paragraphs[0]
    p_para.text = step_badge
    p_para.alignment = PP_ALIGN.CENTER
    p_para.font.size = Pt(9.5)
    p_para.font.bold = True
    p_para.font.color.rgb = RGBColor(255, 255, 255)
    p_para.font.name = "Calibri"

    tb = slide.shapes.add_textbox(Inches(1.95), Inches(0.14), Inches(11.0), Inches(0.88))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p1 = tf.paragraphs[0]
    p1.text = title_text
    p1.font.size = Pt(21)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Calibri"

    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(11)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.font.name = "Calibri"


def add_info_card(slide, left, top, width, height, badge_text, title, bullets, accent_rgb=RGBColor(37, 99, 235), bg_rgb=RGBColor(248, 250, 252), font_size=10.5):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.03
    card.fill.solid()
    card.fill.fore_color.rgb = bg_rgb
    card.line.color.rgb = RGBColor(203, 213, 225)
    card.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.07))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_rgb
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(left + Inches(0.16), top + Inches(0.12), width - Inches(0.32), height - Inches(0.20))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if badge_text:
        p0 = tf.paragraphs[0]
        p0.text = badge_text.upper()
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = accent_rgb
        p0.font.name = "Calibri"
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]

    p_title.text = title
    p_title.font.size = Pt(13)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.font.name = "Calibri"
    p_title.space_after = Pt(5)

    for b in bullets:
        pb = tf.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.size = Pt(font_size)
        pb.font.color.rgb = RGBColor(51, 65, 85)
        pb.font.name = "Calibri"
        pb.space_after = Pt(3.5)


def add_diagram_box(slide, left, top, width, height, header_label, title_text, body_lines, fill_rgb, border_rgb, title_rgb=RGBColor(15, 23, 42)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.adjustments[0] = 0.04
    box.fill.solid()
    box.fill.fore_color.rgb = fill_rgb
    box.line.color.rgb = border_rgb
    box.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.08), width - Inches(0.24), height - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = header_label.upper()
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = border_rgb
    p0.font.name = "Calibri"

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = title_rgb
    p1.font.name = "Calibri"
    p1.space_after = Pt(3)

    for line in body_lines:
        pl = tf.add_paragraph()
        pl.text = line
        pl.font.size = Pt(9.5)
        pl.font.color.rgb = RGBColor(51, 65, 85)
        pl.font.name = "Calibri"
        pl.space_after = Pt(2)


def add_arrow_badge(slide, left, top, width, height, text, bg_rgb=RGBColor(37, 99, 235)):
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arr.fill.solid()
    arr.fill.fore_color.rgb = bg_rgb
    arr.line.fill.background()
    if text:
        tf = arr.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)


def add_screenshot_frame(slide, img_path, left, top, width, height):
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left - Inches(0.05),
        top - Inches(0.05),
        width + Inches(0.10),
        height + Inches(0.10),
    )
    frame.adjustments[0] = 0.015
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor(148, 163, 184)
    frame.line.width = Pt(1.25)

    slide.shapes.add_picture(img_path, left, top, width=width, height=height)


def build_powerpoint_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    s1_img = os.path.join(SCREENSHOT_DIR, "slide_01_control_tower_overview.png")
    s2_img = os.path.join(SCREENSHOT_DIR, "slide_02_gantt_24h_frozen_and_overlaps.png")
    s3_slot_gantt = os.path.join(SCREENSHOT_DIR, "slide_03_slot_swap_gantt.png")
    s4_slot_prop = os.path.join(SCREENSHOT_DIR, "slide_04_slot_swap_proposals.png")
    s5_dock_gantt = os.path.join(SCREENSHOT_DIR, "slide_05_dock_swap_gantt.png")
    s6_dock_prop = os.path.join(SCREENSHOT_DIR, "slide_06_dock_swap_proposals.png")
    s7_modal = os.path.join(SCREENSHOT_DIR, "slide_07_eta_modal_override.png")
    s8_table = os.path.join(SCREENSHOT_DIR, "slide_08_fleet_table_and_logs.png")
    s9_conn = os.path.join(SCREENSHOT_DIR, "slide_09_db_and_external_connections.png")
    s10_trans = os.path.join(SCREENSHOT_DIR, "slide_10_transport_portal.png")
    s11_commit = os.path.join(SCREENSHOT_DIR, "slide_11_approved_committed_schedule.png")

    # =========================================================================
    # SLIDE 1: TITLE & EXECUTIVE HERO SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg1.line.fill.background()

    badge1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.6), Inches(3.9), Inches(0.38))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = RGBColor(37, 99, 235)
    badge1.line.fill.background()
    tf_b1 = badge1.text_frame
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "AUTONOMOUS LOGISTICS AI CONTROL TOWER"
    p_b1.alignment = PP_ALIGN.CENTER
    p_b1.font.size = Pt(10.5)
    p_b1.font.bold = True
    p_b1.font.color.rgb = RGBColor(255, 255, 255)

    tb1 = slide1.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(5.3), Inches(5.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p_t1 = tf1.paragraphs[0]
    p_t1.text = "Autonomous Multi-Agent\nDock & Slot Scheduler"
    p_t1.font.size = Pt(31)
    p_t1.font.bold = True
    p_t1.font.color.rgb = RGBColor(255, 255, 255)
    p_t1.space_after = Pt(10)

    p_sub1 = tf1.add_paragraph()
    p_sub1.text = (
        "Complete Product Architecture, 5-Agent Working Pipeline, "
        "Pairwise Slot Swap & Cross-Dock Swap Scenarios, "
        "24-Hour Scrollable Gantt Chart (30/45/60m Slots), "
        "Overlapping ETA Red Indicator & Carrier Transport Portal."
    )
    p_sub1.font.size = Pt(12.5)
    p_sub1.font.color.rgb = RGBColor(203, 213, 225)
    p_sub1.space_after = Pt(14)

    highlights_s1 = [
        "End-to-End Multi-Agent Architecture & Dual SQLite DB Design",
        "5 Specialized Agents (TruckData, ScheduleData, ETAFinder, SlotChecking, ExecutionComm)",
        "Scenario 1: Pairwise Slot Swap (SLOT_SWAP) on Same Dock",
        "Scenario 2: Cross-Dock Swap (DOCK_SWAP: DOCK-A ↔ DOCK-B)",
        "24-Hour Scrollable Gantt with Frozen Left Labels & 30/45/60m Bars",
        "Overlapping ETA Red Indicator & Standalone /transport Portal",
    ]
    for item in highlights_s1:
        pi = tf1.add_paragraph()
        pi.text = f"✓  {item}"
        pi.font.size = Pt(11)
        pi.font.bold = True
        pi.font.color.rgb = RGBColor(147, 197, 253)
        pi.space_after = Pt(5)

    add_screenshot_frame(slide1, s1_img, Inches(6.05), Inches(0.85), Inches(6.85), Inches(4.11))

    metrics_s1 = [
        ("5 AI Agents", "Autonomous Pipeline"),
        ("Slot & Dock Swap", "Dynamic Re-Allocation"),
        ("24-Hr Gantt", "30 / 45 / 60m Slots"),
        ("3 Connectors", "TMS • WMS • Telematics"),
    ]
    for idx, (m_val, m_lbl) in enumerate(metrics_s1):
        m_left = Inches(6.05 + idx * 1.73)
        m_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left, Inches(5.25), Inches(1.63), Inches(1.35))
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        m_box.line.color.rgb = RGBColor(59, 130, 246)
        tf_m = m_box.text_frame
        tf_m.word_wrap = True
        pm1 = tf_m.paragraphs[0]
        pm1.text = m_val
        pm1.alignment = PP_ALIGN.CENTER
        pm1.font.size = Pt(14)
        pm1.font.bold = True
        pm1.font.color.rgb = RGBColor(96, 165, 250)
        pm2 = tf_m.add_paragraph()
        pm2.text = m_lbl
        pm2.alignment = PP_ALIGN.CENTER
        pm2.font.size = Pt(10)
        pm2.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: SYSTEM ARCHITECTURE DIAGRAM (LAYERED VISUAL BLUEPRINT)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide2, prs, "ARCHITECTURE",
        "End-to-End System & Multi-Database Architecture",
        "4-Tier decoupled architecture: Role-Separated Web Portals, FastAPI Orchestration Layer, 5-Agent Engine, and Pluggable Data Connectors"
    )

    # Tier 1: Left Column — User & Carrier Interfaces
    add_diagram_box(
        slide2, Inches(0.45), Inches(1.38), Inches(2.75), Inches(2.65),
        "TIER 1A • INTERNAL OPS UI",
        "Control Tower Portal (/)",
        [
            "• 24-Hr Scrollable Gantt Chart",
            "• Frozen Track Labels (1. Existing / 2. Proposed)",
            "• Scenario Switcher (Full / Slot / Dock Swap)",
            "• In-Bar Approve/Reject & Bulk Commit",
            "• Fleet Table, Logs & DB Connector UI",
        ],
        RGBColor(239, 246, 255), RGBColor(37, 99, 235)
    )
    add_diagram_box(
        slide2, Inches(0.45), Inches(4.25), Inches(2.75), Inches(2.75),
        "TIER 1B • EXTERNAL CARRIER UI",
        "Transport Portal (/transport)",
        [
            "• Isolated from Main Control Tower",
            "• Card 1: Manual 24-Hr ETA Input",
            "• Card 2: Book New Truck Appointment",
            "• Supports 30m, 45m & 60m Durations",
            "• Instant webhook push to Orchestrator",
        ],
        RGBColor(236, 253, 245), RGBColor(5, 150, 105)
    )

    add_arrow_badge(slide2, Inches(3.28), Inches(2.5), Inches(0.45), Inches(0.35), "")
    add_arrow_badge(slide2, Inches(3.28), Inches(5.4), Inches(0.45), Inches(0.35), "", RGBColor(5, 150, 105))

    # Tier 2: FastAPI REST & State Engine
    add_diagram_box(
        slide2, Inches(3.8), Inches(1.38), Inches(2.45), Inches(5.62),
        "TIER 2 • API GATEWAY",
        "FastAPI Server (server.py)",
        [
            "• GET /api/state (Live Preview)",
            "• POST /api/modify-eta",
            "• POST /api/add-truck",
            "• POST /api/run-agents",
            "• POST /api/proposals/approve",
            "• POST /api/proposals/reject",
            "• POST /api/preset-scenario",
            "  - SCENARIO_FULL_MULTI_AGENT",
            "  - SCENARIO_SLOT_SWAP",
            "  - SCENARIO_DOCK_SWAP",
            "• POST /api/connections/test",
            "• POST /api/connections/save",
            "• POST /api/connections/sync",
            "",
            "Runs Agents 01–04 in non-mutating preview mode on every state poll, and triggers Agent 05 on approval.",
        ],
        RGBColor(248, 250, 252), RGBColor(71, 85, 105)
    )

    add_arrow_badge(slide2, Inches(6.32), Inches(3.95), Inches(0.45), Inches(0.35), "")

    # Tier 3: MultiAgentOrchestrator (5 Agents)
    add_diagram_box(
        slide2, Inches(6.85), Inches(1.38), Inches(3.05), Inches(5.62),
        "TIER 3 • MULTI-AGENT CORE",
        "MultiAgentOrchestrator (agents.py)",
        [
            "01. TruckDataAgent",
            "    Reads fleet manifest, cargo type, weight, duration (30/45/60m) & live ETA.",
            "02. ScheduleDataAgent",
            "    Reads DOCK-A & DOCK-B bay capabilities & 24-hour slot allocations.",
            "03. ETAFinderAgent",
            "    Computes Δ variance, [ETA → ETA+Dur] window & Overlapping ETA Red Alerts.",
            "04. SlotCheckingAgent",
            "    Constraint solver for SLOT_SWAP, DOCK_SWAP, SLOT_TIME_SHIFT & BOOKING.",
            "05. ExecutionCommAgent",
            "    Transactional dual-DB writer & Master Scheduler / Driver dispatcher.",
        ],
        RGBColor(245, 243, 255), RGBColor(124, 58, 237)
    )

    add_arrow_badge(slide2, Inches(9.97), Inches(3.95), Inches(0.45), Inches(0.35), "", RGBColor(124, 58, 237))

    # Tier 4: Dual SQLite Databases & External Connectors
    add_diagram_box(
        slide2, Inches(10.5), Inches(1.38), Inches(2.4), Inches(1.75),
        "TIER 4A • TRUCK DB / TMS",
        "truck_db.sqlite",
        [
            "• Table: trucks",
            "• External: TMS REST API",
            "• External: Postgres URI",
            "• Bound to Agent 01",
        ],
        RGBColor(239, 246, 255), RGBColor(37, 99, 235)
    )
    add_diagram_box(
        slide2, Inches(10.5), Inches(3.31), Inches(2.4), Inches(1.75),
        "TIER 4B • SCHEDULE DB / WMS",
        "schedule_db.sqlite",
        [
            "• Tables: docks, schedule_records, communications",
            "• External: WMS REST/SQL",
            "• Bound to Agent 02 & 05",
        ],
        RGBColor(254, 243, 199), RGBColor(217, 119, 6)
    )
    add_diagram_box(
        slide2, Inches(10.5), Inches(5.25), Inches(2.4), Inches(1.75),
        "TIER 4C • TELEMATICS FEED",
        "Live GPS / ETA Stream",
        [
            "• /api/external-mock/telematics-eta",
            "• Samsara Webhook / API",
            "• Bound to Agent 03",
        ],
        RGBColor(254, 226, 226), RGBColor(220, 38, 38)
    )

    # =========================================================================
    # SLIDE 3: HOW THE 5 AUTONOMOUS AGENTS WORK TOGETHER (DECISION LOGIC)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide3, prs, "AGENTS WORKING",
        "How the 5 Autonomous Agents Work Together — Pipeline & Decision Logic",
        "Step-by-step execution flow from raw truck/schedule ingestion to conflict detection, swap optimization, and atomic DB commit"
    )

    agents_detailed = [
        (
            "AGENT 01",
            "TruckDataAgent",
            "Data Ingestion (TMS)",
            [
                "Connects to truck_db.sqlite or External TMS API.",
                "Loads truck_id, carrier, driver, cargo_type (REFRIGERATED / DRY_VAN / FLATBED), weight_tons, duration_mins (30/45/60m), and current_eta.",
                "Identifies scheduled vs. unbooked inbound trucks.",
            ],
            RGBColor(37, 99, 235),
        ),
        (
            "AGENT 02",
            "ScheduleDataAgent",
            "Dock & Slot State (WMS)",
            [
                "Connects to schedule_db.sqlite or External WMS API.",
                "Loads active docks: DOCK-A (Bay 1: REEFER + DRY_VAN + FLATBED) & DOCK-B (Bay 2: DRY_VAN + FLATBED).",
                "Builds 24-hour occupancy matrix (00:00–24:00) of BOOKED vs. OPEN slots.",
            ],
            RGBColor(124, 58, 237),
        ),
        (
            "AGENT 03",
            "ETAFinderAgent",
            "Variance & Overlap Engine",
            [
                "Computes Δ = ETA_mins − Scheduled_Start_mins.",
                "Flags EARLY (Δ ≤ −30m), DELAYED (Δ ≥ +30m), or UNBOOKED.",
                "Checks same-dock overlap: max(Start_A, Start_B) < min(End_A, End_B) → triggers Red Indicator.",
            ],
            RGBColor(220, 38, 38),
        ),
        (
            "AGENT 04",
            "SlotCheckingAgent",
            "Swap & Slot Optimizer",
            [
                "Priority 1 — Pairwise Slot Swap (SLOT_SWAP): Swaps an early truck with a delayed truck on the same dock.",
                "Priority 2 — Cross-Dock Swap (DOCK_SWAP): Moves a delayed truck to an open compatible bay on the other dock.",
                "Priority 3 — Time Shift / New Booking (30/45/60m).",
            ],
            RGBColor(217, 119, 6),
        ),
        (
            "AGENT 05",
            "ExecutionCommAgent",
            "Transactional Commit",
            [
                "Triggered by single-truck '✓ Approve' or 'Approve All'.",
                "Atomically updates schedule_records in schedule_db.sqlite and trucks in truck_db.sqlite.",
                "Writes audit trail & dispatches alerts to Master Scheduler & Driver.",
            ],
            RGBColor(5, 150, 105),
        ),
    ]

    for idx, (badge_txt, title_txt, sub_txt, bullets, col) in enumerate(agents_detailed):
        left_pos = Inches(0.42 + idx * 2.53)
        add_info_card(
            slide3, left_pos, Inches(1.35), Inches(2.40), Inches(3.85),
            f"{badge_txt} • {sub_txt}", title_txt, bullets, accent_rgb=col, font_size=9.8
        )

    # Bottom Decision Hierarchy Box on Slide 3
    add_info_card(
        slide3, Inches(0.42), Inches(5.38), Inches(12.50), Inches(1.78),
        "MATHEMATICAL CONFLICT & OPTIMIZATION RULES (AGENTS 03 & 04)",
        "How Overlapping ETAs, Slot Swaps, and Cross-Dock Swaps Are Evaluated",
        [
            "1. Live Arrival Window & Overlap Formula: For each truck i with duration D_i ∈ {30, 45, 60} mins, Live Window = [ETA_i, ETA_i + D_i]. Two trucks i, j on the same dock overlap iff max(ETA_i, ETA_j) < min(ETA_i + D_i, ETA_j + D_j).",
            "2. Pairwise Slot Swap (SLOT_SWAP): Selected when Truck A (scheduled at Hour H_1, delayed to H_2) and Truck B (scheduled at Hour H_2, arriving early at H_1) both satisfy cargo compatibility on the dock.",
            "3. Cross-Dock Swap (DOCK_SWAP): Selected when Truck A is delayed to Hour H_target, its current dock D_orig is already occupied at H_target, and alternate dock D_alt supports Truck A's cargo_type and has an OPEN slot at H_target.",
        ],
        accent_rgb=RGBColor(15, 23, 42),
        bg_rgb=RGBColor(241, 245, 249),
        font_size=10.2
    )

    # =========================================================================
    # SLIDE 4: MAIN CONTROL TOWER DASHBOARD (SCREENSHOT 1)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide4, prs, "SCREEN 01",
        "Executive Control Tower & 5-Agent Pipeline Dashboard",
        "Main Scheduler Screen (http://127.0.0.1:8050/) — Scenario Switcher, Condensed KPI Ribbon, and Live 5-Agent Telemetry"
    )
    add_screenshot_frame(slide4, s1_img, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide4, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "CONTROL TOWER HIGHLIGHTS", "Unified Single-Screen Operations",
        [
            "Built-In Scenario Switcher: Instant one-click buttons in the top header for 'Full Mix', 'Slot Swap', and 'Dock Swap (A↔B)'.",
            "Condensed KPI Ribbon: Displays Active Fleet (8), Live ETA Conflicts (5), Pending Agent Proposals (5), and Scheduler Dispatches.",
            "5-Agent Pipeline Stepper: Shows live execution logs and bound data sources (SQLite DB, Telematics API, WMS) for Agents 01–05.",
            "Role Separation: Transport Portal (/transport) is isolated for external carriers with no link exposed on the internal Control Tower.",
        ],
        accent_rgb=RGBColor(37, 99, 235)
    )

    # =========================================================================
    # SLIDE 5: 24-HOUR SCROLLABLE GANTT CHART & FROZEN LABELS (SCREENSHOT 2)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide5, prs, "SCREEN 02",
        "24-Hour Scrollable Gantt Chart, 30/45/60m Slots & Red Overlap Indicator",
        "Synchronized 00:00–24:00 timeline with Frozen Left Track Labels, Variable Duration Bars & Overlapping ETA Red Alerts"
    )
    add_screenshot_frame(slide5, s2_img, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide5, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "24-HR GANTT CAPABILITIES", "Precision Visual Scheduling",
        [
            "24-Hour Scrollable Axis (00:00–24:00): 5760px timeline with :00/:15/:30/:45 ticks and quick-jump buttons (00:00, 06:00, Active 08:00–14:00, 14:00, 18:00).",
            "Frozen Left Track Column: 'Schedule Track', '1. Existing Schedule', and '2. Proposed Schedule' remain locked on the left while scrolling horizontally.",
            "30m, 45m & 60m Appointments: Proportional bar widths (50%, 75%, 100% of hour column) with duration badges.",
            "Overlapping ETA Red Indicator: Red banner + pulsing red slot bars whenever [ETA → ETA+Duration] collides with another truck on the same dock.",
            "Uniform ♦ ETA Pins: Matching red pins (♦ TRK-XXX HH:MM) across both Dock 1 and Dock 2.",
        ],
        accent_rgb=RGBColor(220, 38, 38)
    )

    # =========================================================================
    # SLIDE 6: SCENARIO 1 — PAIRWISE SLOT SWAP (GANTT SCREENSHOT 3)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide6, prs, "SLOT SWAP (1/2)",
        "Scenario 1: Pairwise Slot Swap (SLOT_SWAP) on 24-Hour Gantt Chart",
        "Live Demo ('Slot Swap' button) — Early TRK-103 (ETA 08:05) and Delayed TRK-101 (ETA 10:05) swap appointments on DOCK-A"
    )
    add_screenshot_frame(slide6, s3_slot_gantt, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide6, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "PAIRWISE SLOT SWAP SCENARIO", "Same-Dock Mutual Exchange",
        [
            "Initial Schedule on DOCK-A: TRK-101 (60m) is booked at [08:00–09:00] and TRK-103 (30m) is booked at [10:00–10:30].",
            "Live ETA Disruption: TRK-101 is delayed +125m by highway congestion (Live ETA 10:05), while TRK-103 arrives −115m early (Live ETA 08:05).",
            "Agent 03 & 04 Detection: Agent 03 flags both trucks out-of-window; Agent 04 discovers that TRK-103 can immediately utilize TRK-101's 08:00 bay while TRK-101 takes the 10:00 slot.",
            "Proposed Track 2 Result: TRK-103 moves to DOCK-A [08:00–08:30] ('Slot Swap') and TRK-101 moves to DOCK-A [10:00–11:00] ('Slot Swap') with zero dock idle time.",
        ],
        accent_rgb=RGBColor(124, 58, 237)
    )

    # =========================================================================
    # SLIDE 7: SCENARIO 1 — PAIRWISE SLOT SWAP PROPOSALS & TABLE (SCREENSHOT 4)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide7, prs, "SLOT SWAP (2/2)",
        "Scenario 1: Pairwise Slot Swap — Agent Proposal Cards & Fleet Telemetry",
        "Tab 2 view of SCENARIO_SLOT_SWAP showing exact Agent 04 reasoning, variance metrics, and one-click Approve/Reject controls"
    )
    add_screenshot_frame(slide7, s4_slot_prop, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide7, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "SLOT SWAP AGENT TELEMETRY", "How Agent 04 Validates the Swap",
        [
            "Bidirectional Pairing: Proposal cards explicitly link TRK-101 (DOCK-A 08:00 → DOCK-A 10:00) and TRK-103 (DOCK-A 10:00 → DOCK-A 08:00).",
            "Duration Preservation: TRK-103 retains its 30m appointment duration [08:00–08:30], and TRK-101 retains its 60m duration [10:00–11:00].",
            "Cargo & Bay Check: Both trucks are validated against DOCK-A's supported cargo types (REFRIGERATED, DRY_VAN, FLATBED).",
            "Atomic Approval: Approving either truck commits the paired slot swap in schedule_db.sqlite and truck_db.sqlite.",
        ],
        accent_rgb=RGBColor(124, 58, 237)
    )

    # =========================================================================
    # SLIDE 8: SCENARIO 2 — CROSS-DOCK SWAP (GANTT SCREENSHOT 5)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide8, prs, "DOCK SWAP (1/2)",
        "Scenario 2: Cross-Dock Swap (DOCK_SWAP: DOCK-A ↔ DOCK-B) on Gantt Chart",
        "Live Demo ('Dock Swap (A↔B)' button) — Delayed DOCK-A trucks automatically re-route to open bays on DOCK-B"
    )
    add_screenshot_frame(slide8, s5_dock_gantt, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide8, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "CROSS-DOCK SWAP SCENARIO", "Resolving Same-Dock Collisions",
        [
            "Collision 1 on DOCK-A (11:00): TRK-102 (45m, booked DOCK-A [09:00–09:45]) is delayed to ETA 11:05, colliding with TRK-104 at DOCK-A [11:00].",
            "Cross-Dock Move 1: Because DOCK-A [11:00] is occupied, Agent 04 moves TRK-102 across to open DOCK-B [11:00–11:45] ('Dock Swap').",
            "Collision 2 on DOCK-A (12:00): TRK-104 (45m, booked DOCK-A [11:00–11:45]) is delayed to ETA 12:05, colliding with TRK-105 (60m, ETA 12:00) on DOCK-A [12:00].",
            "Cross-Dock Move 2: Agent 04 flags the Overlapping ETA Red Indicator at 12:05 and moves TRK-104 to open DOCK-B [12:00–12:45] ('Dock Swap').",
        ],
        accent_rgb=RGBColor(217, 119, 6)
    )

    # =========================================================================
    # SLIDE 9: SCENARIO 2 — CROSS-DOCK SWAP PROPOSALS & TABLE (SCREENSHOT 6)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide9, prs, "DOCK SWAP (2/2)",
        "Scenario 2: Cross-Dock Swap — Agent Proposal Cards & Bay Compatibility",
        "Tab 2 view of SCENARIO_DOCK_SWAP showing DOCK_SWAP badges, Overlapping ETA alerts, and cross-bay assignment reasoning"
    )
    add_screenshot_frame(slide9, s6_dock_prop, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide9, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "CROSS-DOCK OPTIMIZER LOGIC", "Constraint-Aware Bay Reallocation",
        [
            "Equipment & Cargo Constraint Check: Before moving TRK-102 (DRY_VAN) and TRK-104 (FLATBED) from DOCK-A to DOCK-B, Agent 04 verifies DOCK-B supports DRY_VAN and FLATBED cargo.",
            "Vacated Slot Recovery: Moving TRK-102 and TRK-104 to DOCK-B frees DOCK-A [09:00] and [11:00] for early or unbooked arrivals.",
            "Zero Detention & Zero Overlap: Eliminates the 12:05–12:50 red overlap between TRK-104 and TRK-105 on DOCK-A.",
            "Driver & Gate Notification: On approval, Agent 05 notifies drivers to proceed to Bay 2 (DOCK-B) instead of Bay 1 (DOCK-A).",
        ],
        accent_rgb=RGBColor(217, 119, 6)
    )

    # =========================================================================
    # SLIDE 10: INTERACTIVE ETA OVERRIDE MODAL & APPROVAL WORKFLOW (SCREENSHOT 7)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide10, prs, "SCREEN 03",
        "Interactive ETA Simulation Modal & In-Bar Approval Controls",
        "Click any truck bar on the Gantt chart to override its 24-hour ETA and trigger instant 5-Agent recalculation"
    )
    add_screenshot_frame(slide10, s7_modal, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide10, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "INTERACTIVE SIMULATION", "What-If ETA Testing & Approval",
        [
            "Click-to-Edit Gantt Bars: Clicking any bar on the Gantt chart opens the native modal dialog for that truck.",
            "24-Hour Time Picker: Override the simulated ETA anywhere between 00:00 and 23:55 and add a telemetry reason note.",
            "Instant Multi-Agent Recalculation: Saving a new ETA immediately runs Agents 01–04 to recompute overlaps, Slot Swaps, and Dock Swaps.",
            "Granular In-Bar Approval: Each proposed bar in Track 2 features '✓ Approve' and '✕' buttons positioned cleanly beside the truck ID.",
        ],
        accent_rgb=RGBColor(37, 99, 235)
    )

    # =========================================================================
    # SLIDE 11: TAB 2 — FULL FLEET TABLE, PROPOSALS & DISPATCH LOGS (SCREENSHOT 8)
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide11, prs, "SCREEN 04",
        "Full Fleet Operations Table, Pending Proposals & Dispatch Feed",
        "Tab 2 ('2. Truck Fleet, ETA Controls & Logs') — Combined Slot Swap, Dock Swap, Time Shift & Unbooked Truck proposals"
    )
    add_screenshot_frame(slide11, s8_table, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide11, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "FLEET & AUDIT WORKSPACE", "Complete Operational Visibility",
        [
            "Pending Schedule Proposals Feed: Displays all 5 active proposals in the Full Mix scenario (SLOT_SWAP, DOCK_SWAP, SLOT_TIME_SHIFT, BOOK_NEW_APPOINTMENT).",
            "Master Scheduler & Driver Communications Log: Real-time audit trail of every committed schedule change and rejection.",
            "Truck Fleet & Live ETA Table: Displays Cargo, Weight, Duration (30m/45m/60m), Original Slot, Current DB Slot, Live ETA window, and Red '⚠ OVERLAPS' badges.",
        ],
        accent_rgb=RGBColor(217, 119, 6)
    )

    # =========================================================================
    # SLIDE 12: TAB 3 — EXTERNAL DB & SYSTEM CONNECTIONS HUB (SCREENSHOT 9)
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide12, prs, "SCREEN 05",
        "Database & External System Connection Links (TMS / WMS / Telematics)",
        "Tab 3 ('3. DB & System Connections') — Pluggable data connectors bound directly to Agents 01, 02, and 03"
    )
    add_screenshot_frame(slide12, s9_conn, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide12, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "EXTERNAL INTEGRATIONS", "Live Connector Architecture",
        [
            "1. Truck DB / TMS Connection (Agent 01): Connect to local SQLite, External TMS SQLite, Mock TMS REST API (/api/external-mock/tms-trucks), or PostgreSQL URI.",
            "2. Schedule DB / WMS Connection (Agent 02): Connect to local Schedule SQLite, External WMS SQLite, Mock WMS REST API, or PostgreSQL URI.",
            "3. Live ETA / Telematics Feed (Agent 03): Connect to Live Telematics REST API (/api/external-mock/telematics-eta), Samsara Webhook, or Truck DB.",
            "Test, Bind & Sync: Includes 'Test Link', 'Save & Bind to Agent', and 'Sync Agent via Link' with live latency & row-count diagnostics.",
        ],
        accent_rgb=RGBColor(8, 145, 178)
    )

    # =========================================================================
    # SLIDE 13: STANDALONE TRANSPORT PORTAL (/transport) (SCREENSHOT 10)
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide13, prs, "SCREEN 06",
        "Standalone Carrier & Driver Transport Portal (/transport)",
        "Dedicated external portal (http://127.0.0.1:8050/transport) for Manual ETA Input & Booking 30/45/60m Appointments"
    )
    add_screenshot_frame(slide13, s10_trans, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide13, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "CARRIER SELF-SERVICE PORTAL", "Clean Two-Card Dispatch UI",
        [
            "Card 01 — Manual ETA Input (Existing Schedule): Carriers select a scheduled truck, view its current slot & duration, and submit an updated 24-hour ETA (00:00–23:55) with quick offset chips.",
            "Card 02 — Book New Truck Schedule / Appointment: Register a new inbound truck (Truck ID, Carrier, Driver, Weight, Cargo Type) and select a 30 min, 45 min, or 60 min appointment duration.",
            "Real-Time Sync to Control Tower: Submissions automatically trigger the 5-Agent pipeline and appear on the Scheduler's Gantt chart.",
        ],
        accent_rgb=RGBColor(5, 150, 105)
    )

    # =========================================================================
    # SLIDE 14: COMMITTED SCHEDULE AFTER MULTI-AGENT APPROVAL (SCREENSHOT 11)
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header_banner(
        slide14, prs, "SCREEN 07",
        "Post-Approval Committed Schedule State on 24-Hour Gantt Chart",
        "After clicking 'Approve All' — All Slot Swaps, Dock Swaps & New Bookings committed to schedule_db.sqlite & truck_db.sqlite"
    )
    add_screenshot_frame(slide14, s11_commit, Inches(0.45), Inches(1.38), Inches(8.6), Inches(5.16))

    add_info_card(
        slide14, Inches(9.25), Inches(1.38), Inches(3.65), Inches(5.72),
        "EXECUTION OUTCOME", "Zero Overlaps & Full Alignment",
        [
            "Pairwise Slot Swap Committed: TRK-103 (30m, early at 08:05) and TRK-101 (60m, delayed to 10:10) swap slots on DOCK-A.",
            "Cross-Dock Swap Committed: TRK-102 (45m, delayed to 11:05) moves from DOCK-A to open DOCK-B [11:00–11:45].",
            "Slot Shift Committed: TRK-104 (45m, delayed to 13:05) shifts to open DOCK-A [13:00–13:45].",
            "New 30m Booking Committed: Unbooked TRK-108 (30m, ETA 09:10) books into vacated DOCK-A [09:00–09:30].",
            "Visual Confirmation: Committed bars turn solid blue ('Approved') and Live ETA Conflicts drop to 0.",
        ],
        accent_rgb=RGBColor(37, 99, 235)
    )

    out_pptx_project = os.path.join(BASE_DIR, "Autonomous_Multi_Agent_Dock_Scheduler_Presentation.pptx")
    out_pptx_artifact = os.path.join(ARTIFACT_DIR, "Autonomous_Multi_Agent_Dock_Scheduler_Presentation.pptx")
    out_pptx_static = os.path.join(STATIC_DIR, "Autonomous_Multi_Agent_Dock_Scheduler_Presentation.pptx")
    out_pptx_downloads = os.path.join(DOWNLOADS_DIR, "Autonomous_Multi_Agent_Dock_Scheduler_Presentation.pptx")

    prs.save(out_pptx_project)
    shutil.copyfile(out_pptx_project, out_pptx_artifact)
    shutil.copyfile(out_pptx_project, out_pptx_static)
    try:
        shutil.copyfile(out_pptx_project, out_pptx_downloads)
    except Exception as e:
        print(f"Note: Could not overwrite Downloads copy (may be open in PowerPoint): {e}")

    print(
        f"Saved 14-Slide PowerPoint presentation to:\n"
        f"  1) {out_pptx_project}\n"
        f"  2) {out_pptx_artifact}\n"
        f"  3) {out_pptx_static}\n"
        f"  4) {out_pptx_downloads}"
    )


if __name__ == "__main__":
    capture_all_product_screenshots()
    build_powerpoint_deck()
